"""
Generates QBG_Plus_Presentation.pptx — a 5-slide deck for the question-bank app.

Slide map (per the brief):
  1. Product introduction
  2. Current challenges with QBG
  3. Fragmented toolchain -> Unified platform
  4. Business impact (cost, TAT, quality)
  5. Future roadmap

Run from the project root:
    python scripts/build_presentation.py
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- design tokens (match the app's purple/indigo palette) ----------
COLOR_BG          = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_TITLE       = RGBColor(0x1E, 0x1B, 0x4B)  # indigo-950
COLOR_ACCENT      = RGBColor(0x63, 0x66, 0xF1)  # indigo-500
COLOR_ACCENT_SOFT = RGBColor(0x81, 0x8C, 0xF8)  # indigo-400
COLOR_VIOLET      = RGBColor(0x8B, 0x5C, 0xF6)
COLOR_BODY        = RGBColor(0x1F, 0x29, 0x37)  # gray-800
COLOR_MUTED       = RGBColor(0x6B, 0x72, 0x80)  # gray-500
COLOR_DIVIDER     = RGBColor(0xE5, 0xE7, 0xEB)  # gray-200
COLOR_CARD_BG     = RGBColor(0xF8, 0xFA, 0xFC)  # near-white card
COLOR_CARD_BORDER = RGBColor(0xE2, 0xE8, 0xF0)
COLOR_RED         = RGBColor(0xEF, 0x44, 0x44)
COLOR_GREEN       = RGBColor(0x10, 0xB9, 0x81)
COLOR_AMBER       = RGBColor(0xF5, 0x9E, 0x0B)

# 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ---------- helpers ----------
def add_text(slide, left, top, width, height, text, *,
             size=18, bold=False, color=COLOR_BODY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font_name="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor

    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_paragraph(text_frame, text, *, size=14, bold=False, color=COLOR_BODY,
                  bullet=False, indent=0, space_before=4):
    p = text_frame.add_paragraph()
    p.level = indent
    if space_before:
        p.space_before = Pt(space_before)
    run = p.add_run()
    if bullet:
        run.text = f"•  {text}"
    else:
        run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return p


def add_rounded_card(slide, left, top, width, height, *,
                     fill=COLOR_CARD_BG, border=COLOR_CARD_BORDER, border_w=0.75):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(border_w)
    shape.shadow.inherit = False
    # remove default text
    if shape.has_text_frame:
        shape.text_frame.text = ""
    return shape


def add_accent_bar(slide, left, top, color=COLOR_ACCENT, height=Pt(4), width=Inches(0.9)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def add_page_chrome(slide, idx, total):
    # Top-left logo dot + product mark
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(0.4), Inches(0.22), Inches(0.22))
    dot.fill.solid()
    dot.fill.fore_color.rgb = COLOR_ACCENT
    dot.line.fill.background()
    dot.shadow.inherit = False
    add_text(slide, Inches(0.8), Inches(0.36), Inches(3), Inches(0.3),
             "QBG+", size=12, bold=True, color=COLOR_TITLE)
    add_text(slide, Inches(1.5), Inches(0.39), Inches(4), Inches(0.3),
             "AI-Powered Question Bank", size=10, color=COLOR_MUTED)

    # Bottom-right page number
    add_text(slide, SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.55),
             Inches(0.9), Inches(0.3),
             f"{idx} / {total}", size=10, color=COLOR_MUTED, align=PP_ALIGN.RIGHT)

    # Thin bottom rule
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  Inches(0.5), SLIDE_H - Inches(0.55),
                                  Inches(0.6), Pt(2))
    rule.fill.solid()
    rule.fill.fore_color.rgb = COLOR_ACCENT
    rule.line.fill.background()
    rule.shadow.inherit = False


def add_slide_heading(slide, eyebrow, title):
    """Adds the standard 'eyebrow + bold title + accent bar' header block."""
    add_text(slide, Inches(0.5), Inches(0.95), Inches(12), Inches(0.4),
             eyebrow.upper(), size=12, bold=True, color=COLOR_ACCENT)
    add_text(slide, Inches(0.5), Inches(1.3), Inches(12), Inches(0.9),
             title, size=34, bold=True, color=COLOR_TITLE)
    add_accent_bar(slide, Inches(0.5), Inches(2.15))


# ---------- slide builders ----------

def slide_1_intro(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_page_chrome(slide, 1, total)

    # Big hero
    add_text(slide, Inches(0.5), Inches(1.0), Inches(12), Inches(0.4),
             "PRODUCT OVERVIEW", size=12, bold=True, color=COLOR_ACCENT)
    add_text(slide, Inches(0.5), Inches(1.4), Inches(12), Inches(1.2),
             "QBG+", size=60, bold=True, color=COLOR_TITLE)
    add_text(slide, Inches(0.5), Inches(2.4), Inches(12), Inches(0.6),
             "An AI-powered, unified platform for managing the JEE / NEET question bank",
             size=20, color=COLOR_BODY)
    add_accent_bar(slide, Inches(0.5), Inches(3.05), width=Inches(1.2), height=Pt(5))

    # Three feature cards across the bottom
    card_top = Inches(3.5)
    card_h   = Inches(3.3)
    cards = [
        ("Manage", "Browse, tag, edit & batch 15 K+ questions with rich metadata, RBAC, and a full audit trail."),
        ("Automate", "One-click test paper generation, AI QC, AI translation, modification, and repeat-check — all built in."),
        ("Review", "Multi-stage QC workflow with named verifiers (1st QC → 2nd QC → UAT) and a per-question edit history."),
    ]
    gap = Inches(0.3)
    card_w = Inches((13.333 - 1.0 - 2 * 0.3) / 3)
    left = Inches(0.5)
    for title, desc in cards:
        add_rounded_card(slide, left, card_top, card_w, card_h)
        add_accent_bar(slide, left + Inches(0.35), card_top + Inches(0.35),
                       width=Inches(0.5), height=Pt(4))
        add_text(slide, left + Inches(0.35), card_top + Inches(0.55),
                 card_w - Inches(0.7), Inches(0.5),
                 title, size=22, bold=True, color=COLOR_TITLE)
        add_text(slide, left + Inches(0.35), card_top + Inches(1.1),
                 card_w - Inches(0.7), Inches(2.0),
                 desc, size=14, color=COLOR_BODY)
        left += card_w + gap


def slide_2_challenges(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_page_chrome(slide, 2, total)
    add_slide_heading(slide, "The problem", "Current challenges with QBG")

    challenges = [
        ("No one-click test paper generation",
         "Selecting questions and assembling a paper is a manual, multi-hour exercise."),
        ("Question repetition within a batch",
         "No safeguard against the same question appearing twice in the same test series."),
        ("Slow content workflow",
         "Filter → copy → paste → format → review — the round trip kills throughput."),
        ("No Single-Digit Integer filter for JEE Advanced",
         "Reviewers manually re-filter to extract the right question type for the paper."),
        ("No saved exam presets",
         "JEE Advanced 2020 / 2021 / 2022 patterns must be rebuilt from scratch every time."),
        ("No AI tooling",
         "QC, translation, modification and repeat-check all happen outside the platform."),
    ]

    # Two-column layout
    col_w = Inches(6.0)
    row_h = Inches(0.95)
    top0 = Inches(2.5)
    for i, (title, desc) in enumerate(challenges):
        col = i % 2
        row = i // 2
        left = Inches(0.5 + col * 6.25)
        top = top0 + row_h * row

        # X-icon
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top + Inches(0.05),
                                      Inches(0.32), Inches(0.32))
        circ.fill.solid()
        circ.fill.fore_color.rgb = RGBColor(0xFE, 0xE2, 0xE2)  # red-100
        circ.line.fill.background()
        circ.shadow.inherit = False
        add_text(slide, left, top + Inches(0.04), Inches(0.32), Inches(0.32),
                 "✕", size=14, bold=True, color=COLOR_RED, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)

        add_text(slide, left + Inches(0.5), top, col_w - Inches(0.5), Inches(0.35),
                 title, size=15, bold=True, color=COLOR_TITLE)
        add_text(slide, left + Inches(0.5), top + Inches(0.4), col_w - Inches(0.5), Inches(0.5),
                 desc, size=12, color=COLOR_MUTED)


def slide_3_unified(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_page_chrome(slide, 3, total)
    add_slide_heading(slide, "The solution", "From fragmented tools to a single platform")

    panel_top = Inches(2.6)
    panel_h   = Inches(4.4)
    panel_w   = Inches(5.7)

    # ── Left panel: Today ──
    add_rounded_card(slide, Inches(0.5), panel_top, panel_w, panel_h,
                     fill=RGBColor(0xFE, 0xF2, 0xF2), border=RGBColor(0xFE, 0xCA, 0xCA))
    add_text(slide, Inches(0.85), panel_top + Inches(0.3), panel_w - Inches(0.7), Inches(0.4),
             "TODAY — FRAGMENTED", size=11, bold=True, color=COLOR_RED)
    add_text(slide, Inches(0.85), panel_top + Inches(0.65), panel_w - Inches(0.7), Inches(0.5),
             "Teams stitching tools together", size=20, bold=True, color=COLOR_TITLE)
    left_items = [
        "Gemini Gems for AI QC and modification",
        "Google Sheets + custom scripts for batch tracking",
        "Manual Word / Doc preparation for the paper",
        "Email / Slack threads for review handoffs",
        "Separate logins, formats, lag at each step",
    ]
    box = slide.shapes.add_textbox(Inches(0.85), panel_top + Inches(1.4),
                                   panel_w - Inches(0.7), panel_h - Inches(1.6))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"•  {left_items[0]}"
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.color.rgb = COLOR_BODY
    for it in left_items[1:]:
        add_paragraph(tf, it, size=14, bullet=True, space_before=8)

    # ── Arrow ──
    arrow_left = Inches(0.5) + panel_w + Inches(0.05)
    arrow_w    = Inches(0.85)
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                   arrow_left, panel_top + Inches(1.9),
                                   arrow_w, Inches(0.5))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = COLOR_ACCENT
    arrow.line.fill.background()
    arrow.shadow.inherit = False

    # ── Right panel: QBG+ ──
    right_left = arrow_left + arrow_w + Inches(0.1)
    add_rounded_card(slide, right_left, panel_top, panel_w, panel_h,
                     fill=RGBColor(0xEC, 0xFD, 0xF5), border=RGBColor(0xA7, 0xF3, 0xD0))
    add_text(slide, right_left + Inches(0.35), panel_top + Inches(0.3),
             panel_w - Inches(0.7), Inches(0.4),
             "WITH QBG+ — UNIFIED", size=11, bold=True, color=COLOR_GREEN)
    add_text(slide, right_left + Inches(0.35), panel_top + Inches(0.65),
             panel_w - Inches(0.7), Inches(0.5),
             "One platform, end-to-end", size=20, bold=True, color=COLOR_TITLE)
    right_items = [
        "AI QC, modification, translation, repeat-check — all built in",
        "Batches, metadata tagging, presets all in-app",
        "QC workflow: Pending → Verified → 2x Verified → UAT",
        "Role-based access + full per-question audit history",
        "Single sign-on, one UI, instant handoffs",
    ]
    box = slide.shapes.add_textbox(right_left + Inches(0.35),
                                   panel_top + Inches(1.4),
                                   panel_w - Inches(0.7), panel_h - Inches(1.6))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"•  {right_items[0]}"
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.color.rgb = COLOR_BODY
    for it in right_items[1:]:
        add_paragraph(tf, it, size=14, bullet=True, space_before=8)


def slide_4_impact(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_page_chrome(slide, 4, total)
    add_slide_heading(slide, "Why it matters", "Business impact")

    # 3 stat cards
    card_top = Inches(2.7)
    card_h   = Inches(3.8)
    gap      = Inches(0.3)
    card_w   = Inches((13.333 - 1.0 - 2 * 0.3) / 3)
    left     = Inches(0.5)

    impacts = [
        ("Cost",
         "Lower content cost",
         "Same content team output multiplied; fewer external AI subscriptions; less duplicated effort across tools.",
         COLOR_ACCENT),
        ("Speed",
         "Faster TAT",
         "Test papers built in minutes instead of days. Presets + one-click generation cut every step that used to be manual.",
         COLOR_VIOLET),
        ("Quality",
         "Fewer errors",
         "AI QC + double verification + UAT + immutable edit history. Every change attributable to a named faculty member.",
         COLOR_GREEN),
    ]

    for eyebrow, title, body, color in impacts:
        add_rounded_card(slide, left, card_top, card_w, card_h)
        # accent stripe top of card
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                        left, card_top, card_w, Pt(6))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = color
        stripe.line.fill.background()
        stripe.shadow.inherit = False

        add_text(slide, left + Inches(0.4), card_top + Inches(0.35),
                 card_w - Inches(0.8), Inches(0.35),
                 eyebrow.upper(), size=12, bold=True, color=color)
        add_text(slide, left + Inches(0.4), card_top + Inches(0.7),
                 card_w - Inches(0.8), Inches(0.7),
                 title, size=22, bold=True, color=COLOR_TITLE)
        add_text(slide, left + Inches(0.4), card_top + Inches(1.55),
                 card_w - Inches(0.8), card_h - Inches(1.7),
                 body, size=14, color=COLOR_BODY)
        left += card_w + gap


def slide_5_roadmap(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_page_chrome(slide, 5, total)
    add_slide_heading(slide, "What's next", "Future roadmap")

    pillars = [
        ("01",
         "Bulk test generation from Sheets",
         "Drop syllabus + paper config into a Google Sheet — the QBG+ backend reads it and emits ready-to-print papers in bulk, with zero manual touch and zero errors."),
        ("02",
         "AI tutor for student doubts",
         "An AI model trained on the question bank + solutions, answering student doubts in real-time with step-by-step explanations grounded in the same content."),
        ("03",
         "Personalised practice papers",
         "Integrate student attempt data (time per question, accuracy patterns) and generate adaptive practice sets tailored to each student's pace and weak areas."),
    ]

    top = Inches(2.6)
    row_h = Inches(1.4)
    for i, (num, title, body) in enumerate(pillars):
        y = top + row_h * i
        # number circle
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                      Inches(0.5), y, Inches(0.95), Inches(0.95))
        circ.fill.solid()
        circ.fill.fore_color.rgb = COLOR_ACCENT
        circ.line.fill.background()
        circ.shadow.inherit = False
        add_text(slide, Inches(0.5), y, Inches(0.95), Inches(0.95),
                 num, size=22, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

        add_text(slide, Inches(1.7), y + Inches(0.1), Inches(11), Inches(0.5),
                 title, size=22, bold=True, color=COLOR_TITLE)
        add_text(slide, Inches(1.7), y + Inches(0.65), Inches(11), Inches(0.7),
                 body, size=13, color=COLOR_BODY)


# ---------- main ----------
def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    builders = [slide_1_intro, slide_2_challenges, slide_3_unified,
                slide_4_impact, slide_5_roadmap]
    total = len(builders)
    for fn in builders:
        fn(prs, total)

    out = Path(__file__).resolve().parents[1] / "QBG_Plus_Presentation.pptx"
    prs.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
